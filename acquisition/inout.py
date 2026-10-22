# -*- coding: utf-8 -*-

import datetime
import csv
import openpyxl
import openpyxl.styles

###############################################################################
def export_csv(sample,transducer,ultrasound,data):
    """Exports data to a CSV file.

    Args:
        sample: Information about the sample.
        transducer: Information about the transducer.
        ultrasound: Ultrasound parameters.
        data: Measured data (AVER, REV, elapsedTime).

    Returns:
        bool: True if export succeeded, False otherwise.
    """
    # Check if sample.filename is valid
    if not hasattr(sample, 'filename') or not sample.filename:
        print("Error: No filename defined in sample.filename.")
        return False

    # Build the full file path
    filename = f"{sample.filename}.csv"
    try:
        with open(filename, 'w', newline='', encoding='utf-8') as csvfile:
            writer = csv.writer(csvfile, delimiter=';')
            print(filename)
            # Write metadata
            writer.writerow(['Information', sample.Info])
            writer.writerow(['Date of experiment', str(datetime.date.today())])
            writer.writerow(['Operator', ''])  # Fill in if needed
            writer.writerow(['Temperature (Celsius degree)', str(sample.temperature)])
            writer.writerow(['Frequency (Hz)', str(ultrasound.frequency)])
            writer.writerow(['Duty cycle (%)', str(ultrasound.DC)])
            writer.writerow(['Amplitude (Vpp)', str(ultrasound.amplitude)])
            writer.writerow(['Number of cycle', str(ultrasound.n_cycle)])
            writer.writerow(['Pulse repetition frequency (Hz)', str(ultrasound.PRF)])
            writer.writerow(['Duration of insonification (s)', str(ultrasound.duration)])
            writer.writerow(['Conversion', str(getattr(ultrasound, 'convertion', 'N/A'))])
            writer.writerow(['Volume fraction (%)', str(sample.Volume_fraction)])
            writer.writerow(['Particle diameter (nm)', str(sample.Diameter)])
            writer.writerow(['Drug concentration (microM)', str(sample.C_Drug)])
            writer.writerow(['Cell type', sample.Cells])

            print(data.AVER)
            print(data.REV)
            print(data.elapsedTime)
            # Write measured data
            if hasattr(data, 'AVER') and hasattr(data, 'REV') and hasattr(data, 'elapsedTime'):
                if len(data.AVER) > 0 and len(data.REV) > 0 and len(data.elapsedTime) > 0:
                    writer.writerow(['Elapsed time (s)', 'AVER (W)', 'REV (W)'])
                    for t, aver, rev in zip(data.elapsedTime, data.AVER, data.REV):
                        writer.writerow([str(t), str(aver), str(rev)])
                else:
                    print("Warning: No data to write (empty arrays).")
            else:
                print("Error: Data arrays not found in data object.")
            print(f"File saved successfully: {filename}")
        return True

    except PermissionError:
        print(f"Error: Permission denied for writing to {filename}.")
        return False
    except Exception as e:
        print(f"Error during CSV export: {e}")
        return False
###############################################################################

###############################################################################
def export_xlsx(sample,transducer,ultrasound,data,iteration):
    r""" This function exports the data in an excel sheet.
    
    :param sample: information about the sample
    :type sample: class

    :param transducer: information about the transducer
    :type transducer: class

    :param ultrasound: information about ultrasound
    :type ultrasound: class

    :param data: information about the data
    :type data: class

    :param iteration: integer
    :type iteration: int      
    
    """
    
    pplusac = data.AVER[1:].mean()*ultrasound.convertion
    pmoinsac = data.REV[1:].mean()*ultrasound.convertion
    pplusmoy = data.AVER[1:].mean()
    pmoinsmoy = data.REV[1:].mean()
    repet = 1/ultrasound.PRF

    print('iteration =',iteration)
    
    if iteration < 1:
        wb = openpyxl.Workbook()
        date_style = openpyxl.styles.NamedStyle(name='mydate',number_format='dd/mm/yy')
        time_style = openpyxl.styles.NamedStyle(name='mytime',number_format='hh:mm:ss')
        print('ok')
        sheet = wb.active
        sheet['A1'] = 'Nanoparticles information'
        sheet['A2'] = 'Name'
        sheet['B2'] = sample.Info
        sheet['A3'] = 'Preparation Date'
        sheet['C3'].style = date_style
        sheet['C3'].number_format = 'dd/mm/yy'
        sheet['C3'] = datetime.date(sample.Preparation_date.year(),sample.Preparation_date.month(),sample.Preparation_date.day())
        sheet['A4'] = 'Diameter (nm) at :'
        sheet['C4'] = 'Day 0'
        sheet['D4'] = sample.Diameter
        sheet['E4'] = 'Day of experiment'
        sheet['G4'].style = date_style
        sheet['G4'] = datetime.date.today()
        sheet['A5'] = 'Drug/Oil name'
        sheet['C5'] = sample.C_Drug
        sheet['A6'] = 'Cells/Solvent'
        sheet['C6'] = sample.Cells
        sheet['A8'] = 'Name of the sample'
        
        sheet['D8'] = 'Date'
        sheet['E8'] = 'Time'
        sheet['F8'] = 'Temp (°C)'
        sheet['G8'] = 'Vol.Fr (%)'
        sheet['H8'] = 'C (µM)'
        sheet['I8'] = 'US dur. (s)'
        sheet['J8'] = 'A (mVpp)'
        sheet['K8'] = 'D.C (%)'
        sheet['L8'] = 'PRF (Hz)'
        sheet['M8'] = 'Psapa (W/cm2)'
        sheet['N8'] = 'Isata (W/cm2)'
        sheet['O8'] = 'Isapa (W/cm2)'
        sheet['P8'] = 'Pavg_ac+(W)'
        sheet['Q8'] = 'Pavg_ac-(W)'
        sheet['R8'] = 'Pavg_el+(W)'
        sheet['S8'] = 'Pavg_el-(W)'
        sheet['T8'] = 'ncycles'
        sheet['U8'] = 'PRP (ms)'
        sheet['V8'].style = time_style

    else:
        wb = openpyxl.load_workbook(sample.filename)
        sheet = wb.active
        
    T = sample.temperature-sample.temperature_zero  
    print(T)
    row = iteration + 9 
    sheet['A'+str(row)] = sample.Info
    sheet['D'+str(row)].style = 'mydate'
    sheet['D'+str(row)] = datetime.date.today()
    sheet['E'+str(row)].style = 'mytime'
    sheet['E'+str(row)] = datetime.datetime.now()
    sheet['F'+str(row)] = sample.temperature
    sheet['G'+str(row)] = sample.Volume_fraction
    sheet['H'+str(row)] = sample.C_Drug
    sheet['I'+str(row)] = ultrasound.duration
    sheet['J'+str(row)] = ultrasound.amplitude
    sheet['K'+str(row)] = ultrasound.DC
    sheet['L'+str(row)] = ultrasound.PRF
    sheet['M'+str(row)] = data.Psapa(T,ultrasound.DC,transducer.efficiency(ultrasound.frequency),transducer.section)
    sheet['N'+str(row)] = data.Isata(transducer.efficiency(ultrasound.frequency),transducer.section)
    sheet['O'+str(row)] = data.Isapa(ultrasound.DC,transducer.efficiency(ultrasound.frequency),transducer.section)
    sheet['P'+str(row)] = pplusac
    sheet['Q'+str(row)] = pmoinsac
    sheet['R'+str(row)] = pplusmoy
    sheet['S'+str(row)] = pmoinsmoy
    sheet['T'+str(row)] = ultrasound.n_cycle
    sheet['U'+str(row)] = repet

    # sauvegarde des puissances instantannées dans une nouvelle feuille
    
    newsheet = wb.create_sheet(title=sample.Info)
    newsheet['A1'] = 'Pavg+'
    newsheet['B1'] = 'Pavg-'
    N = len(data.AVER)
    for i in range(0,N):
        newsheet['A'+str(i+1)] = data.AVER[i]
        newsheet['B'+str(i+1)] = data.REV[i]
        
    wb.save(sample.filename)
    
    return 1
###############################################################################
